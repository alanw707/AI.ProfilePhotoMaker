#!/usr/bin/env python3
"""Build the career market reference snapshot from public BLS releases (ADR 0011).

Usage (needs openpyxl):
  curl -sSLo /tmp/oesm25all.zip https://www.bls.gov/oes/special-requests/oesm25all.zip
  curl -sSLo /tmp/occupation.xlsx https://www.bls.gov/emp/ind-occ-matrix/occupation.xlsx
  python3 scripts/career/build-bls-snapshot.py /tmp/oesm25all.zip /tmp/occupation.xlsx

Writes AI.ProfilePhotoMaker.API/Data/Reference/bls-snapshot.json.gz (deterministic: sorted
keys, fixed gzip mtime) and prints its SHA-256. Values are numbers or a status string,
never a stand-in zero: "*" wage not available, "**" employment not available,
"#" wage at or above the published top code. BLS data are U.S. government works in the
public domain; the citation travels inside the snapshot.
"""
import gzip
import hashlib
import io
import json
import os
import sys
from collections import Counter
import zipfile

import openpyxl

ONET = os.path.join(os.path.dirname(__file__), "..", "..", "AI.ProfilePhotoMaker.API", "Data", "Reference", "onet-snapshot.json.gz")
OUT = os.path.join(os.path.dirname(__file__), "..", "..", "AI.ProfilePhotoMaker.API", "Data", "Reference", "bls-snapshot.json.gz")

# Area types kept: 1 national, 2 state, 4 metropolitan statistical area.
AREA_TYPES = {"1": "national", "2": "state", "4": "metro"}
# Columns kept per wage row, in this order (documented in the snapshot).
WAGE_FIELDS = ["TOT_EMP", "EMP_PRSE", "JOBS_1000", "LOC_QUOTIENT", "A_MEAN", "MEAN_PRSE",
               "A_PCT10", "A_PCT25", "A_MEDIAN", "A_PCT75", "A_PCT90", "H_MEDIAN", "ANNUAL", "HOURLY"]
SENTINELS = {"*", "**", "#"}


def sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def value(v):
    """A number, a BLS status sentinel, or None when the column does not apply."""
    if v is None or v == "" or v == "\u2014":
        # Blank, or the em dash BLS uses for "not applicable" in projections tables.
        return None
    if isinstance(v, (int, float)):
        return round(v, 2) if isinstance(v, float) else v
    s = str(v).strip()
    if s in SENTINELS:
        return s
    if s.upper() == "TRUE":
        return True
    try:
        return round(float(s), 2) if "." in s else int(s)
    except ValueError:
        raise SystemExit(f"Unexpected value {s!r}")


def wages(oews_zip):
    zf = zipfile.ZipFile(oews_zip)
    member = next(n for n in zf.namelist() if n.endswith(".xlsx"))
    wb = openpyxl.load_workbook(io.BytesIO(zf.read(member)), read_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip().upper() for h in next(rows)]
    col = {h: i for i, h in enumerate(header)}
    areas, occupations, data = {}, {}, {}
    for r in rows:
        area_type = str(r[col["AREA_TYPE"]])
        if area_type not in AREA_TYPES or r[col["I_GROUP"]] != "cross-industry":
            continue
        group = r[col["O_GROUP"]]
        if group not in ("detailed", "total"):
            continue
        area, occ = str(r[col["AREA"]]), r[col["OCC_CODE"]]
        areas[area] = {"code": area, "title": r[col["AREA_TITLE"]], "type": AREA_TYPES[area_type], "state": r[col["PRIM_STATE"]]}
        occupations.setdefault(occ, r[col["OCC_TITLE"]])
        data.setdefault(area, {})[occ] = [value(r[col[f]]) for f in WAGE_FIELDS]
    return areas, occupations, data


def projections(xlsx):
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ws = wb["Table 1.2"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[1])
    out = {}
    for r in rows[2:]:
        if not r[1] or r[2] != "Line item":
            continue
        rec = dict(zip(header, r))
        out[r[1]] = {
            "title": str(r[0]).strip(),
            "employment2025Thousands": value(rec["Employment, 2025"]),
            "employment2035Thousands": value(rec["Employment, 2035"]),
            "changeThousands": value(rec["Employment change, numeric, 2025–35"]),
            "changePercent": value(rec["Employment change, percent, 2025–35"]),
            "annualOpeningsThousands": value(rec["Occupational openings, 2025–35 annual average"]),
            "medianAnnualWage2025": value(rec["Median annual wage, dollars, 2025"]),
            "education": rec["Typical education needed for entry"],
            "experience": rec["Work experience in a related occupation"],
            "training": rec["Typical on-the-job training needed to attain competency in the occupation"],
        }
    return out


def crosswalk(onet_codes, available):
    """O*NET-SOC code -> published code: exact SOC detailed, else the SOC broad group."""
    result = {}
    counts = Counter(code[:7] for code in onet_codes)
    for code in onet_codes:
        soc = code[:7]
        broad = soc[:6] + "0"
        if soc in available:
            result[code] = {"code": soc, "match": "shared" if counts[soc] > 1 else "exact"}
        elif broad in available:
            result[code] = {"code": broad, "match": "broad"}
    return result


def main(oews_zip, projections_xlsx):
    onet = json.load(gzip.open(ONET))
    onet_codes = sorted(o["code"] for o in onet["occupations"])

    areas, occupations, data = wages(oews_zip)
    proj = projections(projections_xlsx)
    national = data["99"]

    snapshot = {
        "sources": {
            "oews": {
                "name": "Occupational Employment and Wage Statistics (OEWS), May 2025",
                "publisher": "U.S. Bureau of Labor Statistics",
                "referencePeriod": "2025-05",
                "publishedOn": "2026-05-15",
                "releaseId": "USDL-26-0725",
                "url": "https://www.bls.gov/oes/special-requests/oesm25all.zip",
                "definitionsUrl": "https://www.bls.gov/oes/oes_emp.htm",
                "sourceSha256": sha(oews_zip),
                "license": "Public domain (U.S. government work)",
                "citation": "U.S. Bureau of Labor Statistics, Occupational Employment and Wage Statistics, May 2025 (released May 15, 2026).",
                "wageDefinition": "Straight-time, gross pay of wage and salary workers, excluding overtime, premiums, bonuses, benefits and self-employment income. Annual wages are hourly wages times 2,080 unless only annual wages are published.",
                "employmentDefinition": "Estimated wage and salary employment rounded to the nearest 10; excludes the self-employed.",
                "coverage": "Nonfarm establishments in all 50 states and DC; national, state and metropolitan area cross-industry estimates. Not a count of open jobs.",
                "topCode": {"annual": 239200, "hourly": 115.0},
                "fields": WAGE_FIELDS,
            },
            "projections": {
                "name": "Employment Projections 2025–35, Table 1.2",
                "publisher": "U.S. Bureau of Labor Statistics",
                "referencePeriod": "2025-2035",
                "publishedOn": "2026-08-27",
                "url": "https://www.bls.gov/emp/ind-occ-matrix/occupation.xlsx",
                "definitionsUrl": "https://www.bls.gov/emp/documentation/definitions.htm",
                "sourceSha256": sha(projections_xlsx),
                "license": "Public domain (U.S. government work)",
                "citation": "U.S. Bureau of Labor Statistics, Employment Projections program, 2025–35 (released August 27, 2026).",
                "coverage": "National only. Projected employment and average annual openings from growth and replacement needs; not current vacancies.",
            },
        },
        "areas": [areas[k] for k in sorted(areas)],
        "occupations": {k: occupations[k] for k in sorted(occupations)},
        "wages": {a: {o: data[a][o] for o in sorted(data[a])} for a in sorted(data)},
        "projections": {k: proj[k] for k in sorted(proj)},
        "crosswalk": {
            "oews": crosswalk(onet_codes, national),
            "projections": crosswalk(onet_codes, proj),
        },
    }

    payload = json.dumps(snapshot, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    with open(OUT, "wb") as f:
        with gzip.GzipFile(fileobj=f, mode="wb", mtime=0, filename="", compresslevel=9) as gz:
            gz.write(payload)
    cw = snapshot["crosswalk"]
    print(f"areas {len(areas)}; wage rows {sum(len(v) for v in data.values())}; projections {len(proj)}; "
          f"onet {len(onet_codes)} -> oews {len(cw['oews'])} / projections {len(cw['projections'])}; "
          f"size {os.path.getsize(OUT)}; snapshot sha256 {sha(OUT)}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
